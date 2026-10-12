import csv
import io
import re
from html.parser import HTMLParser
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown", ".csv", ".json", ".html", ".htm", ".xlsx", ".log",
                        ".pptx"}
IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp",
               ".gif": "image/gif", ".bmp": "image/bmp", ".tif": "image/tiff", ".tiff": "image/tiff"}
MIME_TYPES = {".pdf": "application/pdf", ".txt": "text/plain", ".md": "text/markdown", ".markdown": "text/markdown",
              ".csv": "text/csv", ".json": "application/json", ".html": "text/html", ".htm": "text/html",
              ".log": "text/plain",
              ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
              ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
              ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
              **IMAGE_TYPES}


class UnsupportedDocument(ValueError):
    pass


def parse_document(filename: str, data: bytes) -> list[tuple[str, int | None]]:
    """Return a list of (text, page_number) sections."""
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedDocument(f"Unsupported file type '{ext}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")
    if ext == ".pdf":
        sections = _pdf(data)
    elif ext == ".docx":
        sections = _docx(data)
    elif ext in (".html", ".htm"):
        sections = [(html_to_text(_decode(data)), None)]
    elif ext == ".xlsx":
        sections = _xlsx(data)
    elif ext == ".pptx":
        sections = _pptx(data)
    elif ext == ".csv":
        sections = [(_csv(_decode(data)), None)]
    else:
        sections = [(_decode(data), None)]
    sections = [(_normalise(t), p) for t, p in sections]
    sections = [(t, p) for t, p in sections if t]
    if not sections:
        raise UnsupportedDocument("No extractable text found (scanned PDFs need OCR before upload).")
    return sections


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def _normalise(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _pdf(data: bytes) -> list[tuple[str, int | None]]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return [(page.extract_text() or "", i + 1) for i, page in enumerate(reader.pages)]


def _docx(data: bytes) -> list[tuple[str, int | None]]:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = []
    for p in document.paragraphs:
        if not p.text.strip():
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "2"
            parts.append("#" * min(int(level), 6) + " " + p.text.strip())
        else:
            parts.append(p.text)
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [("\n\n".join(parts), None)]


def _xlsx(data: bytes) -> list[tuple[str, int | None]]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for sheet in wb.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            out.append((f"# Sheet: {sheet.title}\n" + "\n".join(rows), None))
    return out


def _pptx(data: bytes) -> list[tuple[str, int | None]]:
    from pptx import Presentation

    out = []
    for i, slide in enumerate(Presentation(io.BytesIO(data)).slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.extend(p.text for p in shape.text_frame.paragraphs if p.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
        if parts:
            out.append((f"# Slide {i}\n" + "\n".join(parts), i))
    return out


def _csv(text: str) -> str:
    return "\n".join(" | ".join(row) for row in csv.reader(io.StringIO(text)))


class _TextExtractor(HTMLParser):
    SKIP = {"script", "style", "noscript", "svg", "nav", "footer", "header", "form"}
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "table"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip += 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self._skip:
            self._skip -= 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def html_to_text(html: str) -> str:
    parser = _TextExtractor()
    try:
        parser.feed(html)
    except Exception:  # malformed markup: keep whatever was extracted
        pass
    return _normalise("".join(parser.parts))
