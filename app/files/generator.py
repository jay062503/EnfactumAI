"""Builds downloadable files from Markdown / tabular content produced by the agent.

Pure functions: (content) -> bytes. No network, no file-system writes, no code execution.
"""
import csv
import html
import io
import json
import re
from dataclasses import dataclass
from pathlib import Path

FORMATS = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv",
    "txt": "text/plain",
    "md": "text/markdown",
    "html": "text/html",
    "json": "application/json",
}
MAX_ROWS = 50_000
_SAFE_NAME = re.compile(r"[^A-Za-z0-9 ._()-]+")
_BOLD = re.compile(r"\*\*(.+?)\*\*")
_INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
_NUMBERED = re.compile(r"^\d+[.)]\s+")


def safe_filename(name: str, fmt: str, fallback: str = "document") -> str:
    stem = _SAFE_NAME.sub("", Path(name or "").stem).strip(" .")[:80] or fallback
    return f"{stem}.{fmt}"


# --------------------------------------------------------------------------- markdown model
@dataclass
class Block:
    kind: str  # h1..h4 | p | ul | ol | table | code | hr
    text: str = ""
    items: list | None = None
    rows: list | None = None


def parse_markdown(md: str) -> list[Block]:
    lines = (md or "").replace("\r\n", "\n").split("\n")
    blocks: list[Block] = []
    i = 0
    para: list[str] = []

    def flush():
        if para:
            blocks.append(Block("p", " ".join(s.strip() for s in para)))
            para.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith("```"):
            flush()
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            blocks.append(Block("code", "\n".join(code)))
        elif not s:
            flush()
        elif m := re.match(r"^(#{1,6})\s+(.*)$", s):
            flush()
            blocks.append(Block(f"h{min(len(m.group(1)), 4)}", m.group(2).strip()))
        elif re.match(r"^(-{3,}|\*{3,}|_{3,})$", s):
            flush()
            blocks.append(Block("hr"))
        elif s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.match(r"^:?-{2,}:?$", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            width = max(len(r) for r in rows) if rows else 0
            blocks.append(Block("table", rows=[r + [""] * (width - len(r)) for r in rows]))
            continue
        elif re.match(r"^[-*+]\s+", s) or _NUMBERED.match(s):
            flush()
            ordered = bool(_NUMBERED.match(s))
            items = []
            while i < len(lines):
                t = lines[i].strip()
                if ordered and _NUMBERED.match(t):
                    items.append(_NUMBERED.sub("", t, count=1))
                elif not ordered and re.match(r"^[-*+]\s+", t):
                    items.append(re.sub(r"^[-*+]\s+", "", t, count=1))
                elif t and items and lines[i].startswith(("  ", "\t")):
                    items[-1] += " " + t
                else:
                    break
                i += 1
            blocks.append(Block("ol" if ordered else "ul", items=items))
            continue
        else:
            para.append(s)
        i += 1
    flush()
    return blocks


def plain(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    return text.replace("**", "").replace("`", "")


# --------------------------------------------------------------------------- PDF
def _fonts():
    try:
        import matplotlib

        base = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
        fonts = {"": base / "DejaVuSans.ttf", "B": base / "DejaVuSans-Bold.ttf", "I": base / "DejaVuSans-Oblique.ttf",
                 "BI": base / "DejaVuSans-BoldOblique.ttf", "mono": base / "DejaVuSansMono.ttf"}
        if all(p.exists() for p in fonts.values()):
            return fonts
    except Exception:
        pass
    return None


def to_pdf(title: str, md: str) -> bytes:
    from fpdf import FPDF, XPos, YPos
    from fpdf.fonts import FontFace

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(True, margin=18)
    pdf.set_margins(18, 18, 18)
    fonts = _fonts()
    if fonts:
        for style in ("", "B", "I", "BI"):
            pdf.add_font("Body", style, str(fonts[style]))
        pdf.add_font("Mono", "", str(fonts["mono"]))
        body, mono = "Body", "Mono"

        def clean(t: str) -> str:
            return t
    else:  # core fonts are latin-1 only
        body, mono = "Helvetica", "Courier"

        def clean(t: str) -> str:
            return t.replace("\u2014", "-").replace("\u2013", "-").replace("\u2019", "'").replace("\u2022", "-") \
                .encode("latin-1", "replace").decode("latin-1")

    pdf.set_title(clean(title))
    pdf.set_creator("Enfactum AI Assistant")
    pdf.add_page()
    width = pdf.epw

    def md_inline(t: str) -> str:
        # fpdf2 markdown: **bold**, __italic__, --underline-- ; neutralise the latter two in normal text
        t = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", t).replace("`", "")
        return clean(t.replace("--", "\u2013").replace("__", "_"))

    def write(w: float, h: float, text: str, **kw):
        pdf.multi_cell(w, h, text, new_x=XPos.LMARGIN, new_y=YPos.NEXT, **kw)

    if title:
        pdf.set_font(body, "B", 18)
        write(width, 9, clean(title))
        pdf.ln(3)
    sizes = {"h1": 16, "h2": 14, "h3": 12.5, "h4": 11.5}
    for b in parse_markdown(md):
        if b.kind in sizes:
            pdf.ln(2)
            pdf.set_font(body, "B", sizes[b.kind])
            write(width, sizes[b.kind] * 0.55, clean(plain(b.text)))
            pdf.ln(1)
        elif b.kind == "p":
            pdf.set_font(body, "", 10.5)
            write(width, 5.6, md_inline(b.text), markdown=True)
            pdf.ln(2)
        elif b.kind in ("ul", "ol"):
            pdf.set_font(body, "", 10.5)
            for n, item in enumerate(b.items or [], start=1):
                bullet = f"{n}." if b.kind == "ol" else ("\u2022" if fonts else "-")
                pdf.set_x(pdf.l_margin + 3)
                pdf.cell(6, 5.6, bullet, new_x=XPos.RIGHT, new_y=YPos.TOP)
                write(width - 9, 5.6, md_inline(item), markdown=True)
            pdf.ln(2)
        elif b.kind == "code":
            pdf.set_font(mono, "", 9)
            pdf.set_fill_color(244, 245, 247)
            write(width, 4.8, clean(b.text), fill=True)
            pdf.ln(2)
        elif b.kind == "table" and b.rows:
            pdf.set_font(body, "", 9)
            with pdf.table(width=width, line_height=5.2, text_align="LEFT",
                           headings_style=FontFace(emphasis="BOLD", fill_color=(235, 238, 245))) as table:
                for r in b.rows:
                    row = table.row()
                    for c in r:
                        row.cell(clean(plain(c)))
            pdf.ln(3)
        elif b.kind == "hr":
            y = pdf.get_y() + 2
            pdf.set_draw_color(210, 214, 220)
            pdf.line(pdf.l_margin, y, pdf.l_margin + width, y)
            pdf.ln(5)
    return bytes(pdf.output())


# --------------------------------------------------------------------------- DOCX
def _docx_runs(par, text: str):
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            par.add_run(plain(part))


def to_docx(title: str, md: str) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.core_properties.title = title
    doc.core_properties.author = "Enfactum AI Assistant"
    if title:
        doc.add_heading(title, 0)
    for b in parse_markdown(md):
        if b.kind.startswith("h"):
            doc.add_heading(plain(b.text), int(b.kind[1]))
        elif b.kind == "p":
            _docx_runs(doc.add_paragraph(), b.text)
        elif b.kind in ("ul", "ol"):
            for item in b.items or []:
                _docx_runs(doc.add_paragraph(style="List Number" if b.kind == "ol" else "List Bullet"), item)
        elif b.kind == "code":
            run = doc.add_paragraph().add_run(b.text)
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif b.kind == "table" and b.rows:
            table = doc.add_table(rows=len(b.rows), cols=len(b.rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(b.rows):
                for c, value in enumerate(row):
                    cell = table.cell(r, c)
                    cell.text = plain(value)
                    if r == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            doc.add_paragraph()
        elif b.kind == "hr":
            doc.add_paragraph("\u2014" * 30)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# --------------------------------------------------------------------------- PPTX
def to_pptx(title: str, md: str) -> bytes:
    """`## Heading` starts a new slide; bullets / paragraphs become slide points; tables become slide tables."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    first = prs.slides.add_slide(prs.slide_layouts[0])
    first.shapes.title.text = title or "Presentation"
    first.placeholders[1].text = "Prepared with Enfactum AI Assistant"

    slides: list[tuple[str, list[Block]]] = []
    for b in parse_markdown(md):
        if b.kind in ("h1", "h2") or not slides:
            slides.append((plain(b.text) if b.kind in ("h1", "h2") else title, []))
            if b.kind in ("h1", "h2"):
                continue
        slides[-1][1].append(b)

    for heading, blocks in slides:
        table = next((b for b in blocks if b.kind == "table" and b.rows), None)
        slide = prs.slides.add_slide(prs.slide_layouts[5 if table else 1])
        slide.shapes.title.text = heading
        if table:
            rows, cols = len(table.rows), len(table.rows[0])
            shape = slide.shapes.add_table(rows, cols, Inches(0.6), Inches(1.6), Inches(12.1),
                                           Inches(min(0.4 * rows, 5.4)))
            for r, row in enumerate(table.rows):
                for c, value in enumerate(row):
                    cell = shape.table.cell(r, c)
                    cell.text = plain(value)
                    cell.text_frame.paragraphs[0].font.size = Pt(12)
            continue
        frame = slide.placeholders[1].text_frame
        frame.clear()
        points = []
        for b in blocks:
            if b.kind in ("ul", "ol"):
                points.extend((plain(i), 0) for i in b.items or [])
            elif b.kind in ("p", "h3", "h4", "code"):
                points.append((plain(b.text), 0 if b.kind != "code" else 1))
        for n, (text, level) in enumerate(points or [("", 0)]):
            p = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
            p.text = text
            p.level = level
            p.font.size = Pt(20 if len(points) <= 6 else 16)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


# --------------------------------------------------------------------------- tabular
def _sheets(columns, rows, sheets) -> list[dict]:
    if sheets:
        return [{"name": str(s.get("name") or f"Sheet{i + 1}")[:31], "columns": s.get("columns") or [],
                 "rows": (s.get("rows") or [])[:MAX_ROWS]} for i, s in enumerate(sheets)]
    return [{"name": "Sheet1", "columns": columns or [], "rows": (rows or [])[:MAX_ROWS]}]


def _table_from_md(md: str) -> tuple[list, list]:
    for b in parse_markdown(md):
        if b.kind == "table" and b.rows:
            return b.rows[0], b.rows[1:]
    return [], []


_NUMBER = re.compile(r"[-+]?\d+(\.\d+)?")


def _formula_like(v) -> bool:
    """Cells starting with = + - @ are executed by Excel/Sheets (CSV/formula injection) unless plain numbers."""
    return isinstance(v, str) and bool(v) and v[0] in "=+-@" and not _NUMBER.fullmatch(v.strip())


def _cell(v):
    if isinstance(v, str):
        s = v.strip().replace(",", "")
        if _NUMBER.fullmatch(s) and len(s) < 16:
            return float(s) if "." in s else int(s)
        if _formula_like(v):
            return "'" + v
    return v


def to_xlsx(title: str, columns=None, rows=None, sheets=None, md: str = "") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    if not columns and not rows and not sheets and md:
        columns, rows = _table_from_md(md)
    wb = Workbook()
    wb.remove(wb.active)
    for spec in _sheets(columns, rows, sheets):
        ws = wb.create_sheet(_SAFE_NAME.sub("", spec["name"]) or "Sheet")
        if spec["columns"]:
            ws.append([str(c) for c in spec["columns"]])
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2F5BEA")
                cell.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A2"
        for row in spec["rows"]:
            ws.append([_cell(v) for v in (row if isinstance(row, list) else [row])])
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
        if spec["columns"] and spec["rows"]:
            ws.auto_filter.ref = ws.dimensions
    wb.properties.title = title
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_csv(columns=None, rows=None, md: str = "") -> bytes:
    if not columns and not rows and md:
        columns, rows = _table_from_md(md)
    out = io.StringIO()
    w = csv.writer(out)
    if columns:
        w.writerow(columns)
    for row in (rows or [])[:MAX_ROWS]:
        w.writerow([("'" + v) if _formula_like(v) else v
                    for v in (row if isinstance(row, list) else [row])])
    return ("\ufeff" + out.getvalue()).encode("utf-8")


# --------------------------------------------------------------------------- text formats
def to_html(title: str, md: str) -> bytes:
    def inline(t: str) -> str:
        t = html.escape(t)
        t = _BOLD.sub(r"<strong>\1</strong>", t)
        t = re.sub(r"`([^`]+)`", r"<code>\1</code>", t)
        return re.sub(r"\[([^\]]+)\]\((https?://[^)\s]+)\)", r'<a href="\2">\1</a>', t)

    parts = []
    for b in parse_markdown(md):
        if b.kind.startswith("h"):
            parts.append(f"<{b.kind}>{inline(b.text)}</{b.kind}>")
        elif b.kind == "p":
            parts.append(f"<p>{inline(b.text)}</p>")
        elif b.kind in ("ul", "ol"):
            parts.append(f"<{b.kind}>" + "".join(f"<li>{inline(i)}</li>" for i in b.items or []) + f"</{b.kind}>")
        elif b.kind == "code":
            parts.append(f"<pre><code>{html.escape(b.text)}</code></pre>")
        elif b.kind == "table" and b.rows:
            head = "".join(f"<th>{inline(c)}</th>" for c in b.rows[0])
            body = "".join("<tr>" + "".join(f"<td>{inline(c)}</td>" for c in r) + "</tr>" for r in b.rows[1:])
            parts.append(f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>")
        elif b.kind == "hr":
            parts.append("<hr>")
    style = ("body{font-family:system-ui,-apple-system,Segoe UI,Roboto,sans-serif;max-width:860px;margin:40px auto;"
             "padding:0 20px;line-height:1.6;color:#1f2430}table{border-collapse:collapse;width:100%;margin:12px 0}"
             "th,td{border:1px solid #d9dde5;padding:6px 10px;text-align:left}th{background:#eef1f7}"
             "pre{background:#f4f5f7;padding:12px;border-radius:8px;overflow:auto}code{font-family:ui-monospace,monospace}")
    return (f"<!doctype html><html><head><meta charset=\"utf-8\"><title>{html.escape(title)}</title>"
            f"<style>{style}</style></head><body><h1>{html.escape(title)}</h1>{''.join(parts)}</body></html>"
            ).encode("utf-8")


def to_txt(title: str, md: str) -> bytes:
    lines = [title, "=" * len(title), ""] if title else []
    for b in parse_markdown(md):
        if b.kind.startswith("h"):
            lines += ["", plain(b.text).upper() if b.kind in ("h1", "h2") else plain(b.text), ""]
        elif b.kind == "p":
            lines += [plain(b.text), ""]
        elif b.kind in ("ul", "ol"):
            lines += [f"{f'{n}.' if b.kind == 'ol' else '-'} {plain(i)}" for n, i in enumerate(b.items or [], 1)]
            lines.append("")
        elif b.kind == "code":
            lines += [b.text, ""]
        elif b.kind == "table" and b.rows:
            widths = [max(len(plain(r[c])) for r in b.rows) for c in range(len(b.rows[0]))]
            for n, r in enumerate(b.rows):
                lines.append("  ".join(plain(v).ljust(widths[c]) for c, v in enumerate(r)))
                if n == 0:
                    lines.append("  ".join("-" * w for w in widths))
            lines.append("")
        elif b.kind == "hr":
            lines += ["-" * 40, ""]
    return "\n".join(lines).strip().encode("utf-8") + b"\n"


def to_json(content: str, columns=None, rows=None) -> bytes:
    if columns and rows:
        data = [dict(zip(columns, r)) for r in rows[:MAX_ROWS]]
    else:
        try:
            data = json.loads(content)
        except ValueError:
            data = {"content": content}
    return json.dumps(data, indent=2, ensure_ascii=False, default=str).encode("utf-8")


def build_file(fmt: str, *, title: str, content: str = "", columns=None, rows=None, sheets=None) -> bytes:
    fmt = fmt.lower()
    if fmt == "pdf":
        return to_pdf(title, content)
    if fmt == "docx":
        return to_docx(title, content)
    if fmt == "pptx":
        return to_pptx(title, content)
    if fmt == "xlsx":
        return to_xlsx(title, columns, rows, sheets, content)
    if fmt == "csv":
        return to_csv(columns, rows, content)
    if fmt == "html":
        return to_html(title, content)
    if fmt == "txt":
        return to_txt(title, content)
    if fmt == "md":
        return ((f"# {title}\n\n" if title and not content.lstrip().startswith("#") else "") + content).encode("utf-8")
    if fmt == "json":
        return to_json(content, columns, rows)
    raise ValueError(f"Unsupported format: {fmt}")


# --------------------------------------------------------------------------- charts
CHART_TYPES = ("bar", "barh", "line", "area", "pie", "donut", "scatter", "stacked_bar")
PALETTE = ["#2F5BEA", "#14B8A6", "#F59E0B", "#EF4444", "#8B5CF6", "#10B981", "#EC4899", "#64748B"]


def build_chart(*, chart_type: str, title: str, labels: list, series: list[dict], x_label: str = "",
                y_label: str = "") -> bytes:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    labels = [str(x) for x in labels][:500]
    series = [{"name": str(s.get("name", f"Series {i + 1}")),
               "values": [float(v) if v is not None else 0.0 for v in (s.get("values") or [])][:len(labels) or 500]}
              for i, s in enumerate(series[:8])]
    fig, ax = plt.subplots(figsize=(10, 5.6), dpi=150)
    fig.patch.set_facecolor("white")
    if chart_type in ("pie", "donut"):
        vals = series[0]["values"] if series else []
        ax.pie(vals, labels=labels[:len(vals)], colors=PALETTE * 10, autopct="%1.1f%%", startangle=90,
               wedgeprops={"width": 0.45} if chart_type == "donut" else None, textprops={"fontsize": 9})
        ax.axis("equal")
    else:
        import numpy as np

        x = np.arange(len(labels))
        n = max(len(series), 1)
        bottom = np.zeros(len(labels))
        for i, s in enumerate(series):
            vals = np.array((s["values"] + [0.0] * len(labels))[:len(labels)])
            color = PALETTE[i % len(PALETTE)]
            if chart_type == "bar":
                w = 0.8 / n
                ax.bar(x - 0.4 + w * (i + 0.5), vals, w, label=s["name"], color=color)
            elif chart_type == "stacked_bar":
                ax.bar(x, vals, 0.6, bottom=bottom, label=s["name"], color=color)
                bottom += vals
            elif chart_type == "barh":
                h = 0.8 / n
                ax.barh(x - 0.4 + h * (i + 0.5), vals, h, label=s["name"], color=color)
            elif chart_type == "area":
                ax.fill_between(x, vals, alpha=0.25, color=color)
                ax.plot(x, vals, label=s["name"], color=color, linewidth=2)
            elif chart_type == "scatter":
                ax.scatter(x, vals, label=s["name"], color=color)
            else:
                ax.plot(x, vals, marker="o", label=s["name"], color=color, linewidth=2)
        if chart_type == "barh":
            ax.set_yticks(x, labels)
        else:
            ax.set_xticks(x, labels, rotation=30 if max((len(lb) for lb in labels), default=0) > 6 else 0,
                          ha="right" if max((len(lb) for lb in labels), default=0) > 6 else "center")
        ax.set_xlabel(x_label)
        ax.set_ylabel(y_label)
        ax.grid(axis="x" if chart_type == "barh" else "y", alpha=0.3)
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)
        if len(series) > 1:
            ax.legend(frameon=False)
    ax.set_title(title, fontsize=13, fontweight="bold", pad=12)
    fig.tight_layout()
    out = io.BytesIO()
    fig.savefig(out, format="png")
    plt.close(fig)
    return out.getvalue()
